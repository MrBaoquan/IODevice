# -*- coding: utf-8 -*-
"""生成节点一汇报 PPT —— 集成式动感平台播控系统
汇报结构：计划目标 → 已完成工作 → 执行过程 → 难点突破 → 下一步计划 → 人员分工与倒排计划
节点描述严格按立项报告
v5: 11 项修改 — 封面简化/节点标签/影片播控平台前置/单页展示/厂家对接整合/多人并行
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os, sys

# 解决 Windows GBK 编码问题
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "节点一_进度汇报.pptx")

# 图片路径
IMG_EDITOR = os.path.join(SCRIPT_DIR, "动作编辑平台-1.png")
IMG_TESTTOOL = os.path.join(SCRIPT_DIR, "上海飞荇客测试工具-gui.png")
IMG_CHAT1 = os.path.join(SCRIPT_DIR, "d55fe844d2d541439f2452dc2a7f7fdd.png")
IMG_CHAT2 = os.path.join(SCRIPT_DIR, "获取到飞荇客协议-1.png")
IMG_CHAT3 = os.path.join(SCRIPT_DIR, "获取到飞荇客协议-2.png")
IMG_TIMELINE = os.path.join(SCRIPT_DIR, "..", "..", "image", "动感平台_进度汇报", "1774685442044.png")

# 影片播控平台界面原型
OWNER_DIR = os.path.join(SCRIPT_DIR, "..", "..", "Pencil", "exports")
IMG_OW_STANDBY  = os.path.join(OWNER_DIR, "1-待机界面.png")
IMG_OW_SELECT   = os.path.join(OWNER_DIR, "2-影片选择界面.png")
IMG_OW_PLAY     = os.path.join(OWNER_DIR, "3-影片播放界面.png")
IMG_OW_EMERG    = os.path.join(OWNER_DIR, "4-紧急处理界面.png")
IMG_OW_END      = os.path.join(OWNER_DIR, "5-影片结束界面.png")
IMG_OW_MONITOR  = os.path.join(OWNER_DIR, "6-现场监控界面.png")

# ── 播控系统配色方案 (深蓝 + 金色影院风格) ──
C_DARK    = RGBColor(0x0A, 0x19, 0x29)   # 深邃夜蓝 - 背景
C_PRIMARY = RGBColor(0x14, 0x5C, 0xA8)   # 播控蓝 - 主色调
C_ACCENT  = RGBColor(0xD4, 0xA0, 0x2E)   # 影院金 - 强调/高亮
C_TEAL    = RGBColor(0x00, 0x89, 0x7B)   # 科技青 - 辅助正面色
C_ORANGE  = RGBColor(0xD1, 0x6B, 0x1E)   # 暖橙 - 警示/难点
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT_BG= RGBColor(0xEE, 0xF2, 0xF7)   # 浅灰蓝背底
C_GRAY    = RGBColor(0x55, 0x5E, 0x68)
C_LGRAY   = RGBColor(0x8C, 0x95, 0x9F)
C_TBL_HDR = RGBColor(0x0E, 0x3D, 0x6B)   # 表头深蓝
C_TBL_R1  = RGBColor(0xE5, 0xEE, 0xF7)   # 表格奇行
C_TBL_R2  = RGBColor(0xFF, 0xFF, 0xFF)
C_OK_BG   = RGBColor(0xC8, 0xE6, 0xC9)   # 已完成绿底(更明显)
C_OK_TEXT = RGBColor(0x1B, 0x5E, 0x20)   # 已完成深绿文字
C_HL_BG   = RGBColor(0xFD, 0xF4, 0xE0)   # 核心高亮底色（金色淡底）
C_CUR_BG  = RGBColor(0xFF, 0xF3, 0xCD)   # 当前节点高亮行
# 新增: 厂家对接事件颜色
C_RED_BG    = RGBColor(0xFD, 0xE8, 0xE8)   # 不顺利事件浅红底
C_RED_TEXT  = RGBColor(0xC6, 0x28, 0x28)   # 红色文字
C_GREEN_TEXT= RGBColor(0x2E, 0x7D, 0x32)   # 绿色文字

# 节点标签统一定义 (Change #11)
NODE_LABELS = {
    1: "节点一(架构原型)",
    2: "节点二(核心开发)",
    3: "节点三(接口整合)",
    4: "节点四(测试优化)",
    5: "节点五(产品验收)",
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


# ═══════════ 工具函数 ═══════════

def bg(slide, color):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = color

def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    return s

def rrect(slide, l, t, w, h, color, border=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    if border:
        s.line.color.rgb = border; s.line.width = Pt(2)
    else:
        s.line.fill.background()
    return s

def badge(slide, l, t, w, h, text, bg_color, sz=14):
    """宽序号徽章，避免数字换行"""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = bg_color; s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(sz); p.font.color.rgb = C_WHITE
    p.font.bold = True; p.font.name = "微软雅黑"; p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    tf.paragraphs[0].space_after = Pt(0)
    s.text_frame.margin_top = Pt(2)
    s.text_frame.margin_bottom = Pt(2)
    return s

def txt(slide, l, t, w, h, text, sz=18, color=C_DARK, bold=False,
        align=PP_ALIGN.LEFT, font="微软雅黑"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(sz); p.font.color.rgb = color
    p.font.bold = bold; p.font.name = font; p.alignment = align
    return tb

def mtxt(slide, l, t, w, h, lines, sz=16, color=C_DARK, font="微软雅黑"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (text, is_bold, c) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text; p.font.size = Pt(sz)
        p.font.color.rgb = c if c else color
        p.font.bold = is_bold; p.font.name = font
        p.space_after = Pt(sz * 0.3)
    return tb

def tbl_ex(slide, l, t, w, data, col_ws, highlight_row=None, row_colors=None):
    """表格，支持高亮行和自定义行背景色 row_colors={row_idx: RGBColor}"""
    nr, nc = len(data), len(data[0])
    ts = slide.shapes.add_table(nr, nc, l, t, w, Inches(0.4) * nr)
    table = ts.table
    for i, cw in enumerate(col_ws): table.columns[i].width = cw
    for ri, row in enumerate(data):
        for ci, ct in enumerate(row):
            cell = table.cell(ri, ci); cell.text = ""
            p = cell.text_frame.paragraphs[0]; p.text = str(ct); p.font.name = "微软雅黑"
            if ri == 0:
                p.font.size = Pt(12); p.font.color.rgb = C_WHITE; p.font.bold = True
                p.alignment = PP_ALIGN.CENTER; cell.fill.solid(); cell.fill.fore_color.rgb = C_TBL_HDR
            elif row_colors and ri in row_colors:
                rc = row_colors[ri]
                p.font.size = Pt(12); p.font.bold = True
                p.font.color.rgb = rc.get('text', C_DARK) if isinstance(rc, dict) else C_DARK
                cell.fill.solid(); cell.fill.fore_color.rgb = rc.get('bg', rc) if isinstance(rc, dict) else rc
            elif highlight_row is not None and ri == highlight_row:
                p.font.size = Pt(12); p.font.color.rgb = C_DARK; p.font.bold = True
                cell.fill.solid(); cell.fill.fore_color.rgb = C_CUR_BG
            else:
                p.font.size = Pt(11); p.font.color.rgb = C_DARK; cell.fill.solid()
                cell.fill.fore_color.rgb = C_TBL_R1 if ri % 2 == 1 else C_TBL_R2
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Pt(6); cell.margin_right = Pt(6)
            cell.margin_top = Pt(4); cell.margin_bottom = Pt(4)
    return ts

def tbl(slide, l, t, w, data, col_ws):
    return tbl_ex(slide, l, t, w, data, col_ws)

def title_bar(slide, text):
    rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.0), C_DARK)
    rect(slide, Inches(0), Inches(0.96), SLIDE_W, Inches(0.04), C_ACCENT)  # 金色装饰线
    txt(slide, Inches(0.6), Inches(0.15), Inches(12), Inches(0.7),
        text, sz=30, color=C_WHITE, bold=True)

def add_img(slide, path, l, t, w=None, h=None):
    if not os.path.exists(path): return None
    if w and h: return slide.shapes.add_picture(path, l, t, w, h)
    if w: return slide.shapes.add_picture(path, l, t, width=w)
    return slide.shapes.add_picture(path, l, t)


# ═══════════════════════════════════════════
# Slide 1: 封面 (Change #1: 简洁风格，白色背景，居中布局)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)

# 顶部装饰线
rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.08), C_PRIMARY)

# 居中标题
txt(s, Inches(0), Inches(2.2), SLIDE_W, Inches(1.0),
    "集成式动感平台播控系统", sz=48, color=C_PRIMARY, bold=True,
    align=PP_ALIGN.CENTER)
txt(s, Inches(0), Inches(3.3), SLIDE_W, Inches(0.6),
    "第一阶段 · 进度汇报", sz=28, color=C_DARK, bold=True,
    align=PP_ALIGN.CENTER)

# 分隔线
rect(s, Inches(5.6), Inches(4.15), Inches(2.1), Inches(0.03), C_ACCENT)

# 软件技术中心
txt(s, Inches(0), Inches(4.6), SLIDE_W, Inches(0.5),
    "软件技术中心", sz=22, color=C_GRAY, align=PP_ALIGN.CENTER)

# 日期
txt(s, Inches(0), Inches(5.2), SLIDE_W, Inches(0.4),
    "2026 年 3 月 31 日", sz=18, color=C_LGRAY, align=PP_ALIGN.CENTER)

# 汇报人
txt(s, Inches(0), Inches(5.8), SLIDE_W, Inches(0.4),
    "汇报人：xxx", sz=16, color=C_LGRAY, align=PP_ALIGN.CENTER)

# 底部装饰线
rect(s, Inches(0), Inches(7.38), SLIDE_W, Inches(0.08), C_PRIMARY)


# ═══════════════════════════════════════════
# Slide 2: 项目进度规划 (Change #2: 节点一绿色已完成, Change #11: 节点标签)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "项目进度规划（立项报告）")

plan_data = [
    ["节点", "时间节点", "进度安排", "状态"],
    [NODE_LABELS[1], "2026 年 3 月 31 日前", "产品原型设计及系统架构方案完成", "已按计划完成"],
    [NODE_LABELS[2], "2026 年 4 月 30 日前", "核心播控系统开发与基础联调完成", ""],
    [NODE_LABELS[3], "2026 年 5 月 15 日前", "完成设备接口整合及动作编排优化", ""],
    [NODE_LABELS[4], "2026 年 5 月 31 日前", "系统整体测试优化完成，达到上线标准", ""],
    [NODE_LABELS[5], "2026 年 6 月 30 日前", "产品验收完成", ""],
]
tbl_ex(s, Inches(0.6), Inches(1.3), Inches(12), plan_data,
       [Inches(2.0), Inches(2.8), Inches(5.4), Inches(1.8)],
       row_colors={1: {'bg': C_OK_BG, 'text': C_OK_TEXT}})

# 当前节点指示条
rrect(s, Inches(0.6), Inches(3.95), Inches(12), Inches(0.5), C_ACCENT)
txt(s, Inches(0.6), Inches(3.97), Inches(12), Inches(0.5),
    f"当前汇报：{NODE_LABELS[1]} — 产品原型设计及系统架构方案完成（截止 3 月 31 日）",
    sz=15, color=C_DARK, bold=True, align=PP_ALIGN.CENTER)

# 本项目需交付的主要产品 (Change #3: 流程控制软件 → 影片播控平台)
txt(s, Inches(0.6), Inches(4.75), Inches(12), Inches(0.4),
    "项目需交付的主要产品", sz=18, color=C_DARK, bold=True)

products = [
    ("影片播控平台",
     "提供给影院现场使用，控制视频播放与动感平台联动，全流程影片播控管理",
     C_ACCENT),
    ("动作编辑平台",
     "可视化时间轴编辑工具，用于编排动感座椅运动数据，支持视频同步预览",
     C_PRIMARY),
]
for i, (title, desc, color) in enumerate(products):
    x = Inches(0.6) + Inches(6.0) * i
    y = Inches(5.3)
    rrect(s, x, y, Inches(5.7), Inches(1.8), C_LIGHT_BG)
    rect(s, x, y + Inches(0.15), Inches(0.06), Inches(1.4), color)
    badge(s, x + Inches(0.2), y + Inches(0.2), Inches(0.6), Inches(0.35),
          str(i + 1), color, sz=14)
    txt(s, x + Inches(0.95), y + Inches(0.15), Inches(4.5), Inches(0.4),
        title, sz=17, color=C_DARK, bold=True)
    txt(s, x + Inches(0.25), y + Inches(0.65), Inches(5.2), Inches(1.0),
        desc, sz=13, color=C_GRAY)


# ═══════════════════════════════════════════
# Slide 3: 本阶段计划目标分解  (Change #3: 术语更新)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "一、本阶段计划目标 — 产品原型设计及系统架构方案完成")

# 三大分类标题
txt(s, Inches(0.6), Inches(1.15), Inches(4), Inches(0.42),
    "核心工作", sz=20, color=C_ACCENT, bold=True)
txt(s, Inches(0.6), Inches(1.48), Inches(4), Inches(0.3),
    "系统架构原型 + 平台协议沟通", sz=12, color=C_LGRAY)

# 核心工作卡片
core_items = [
    ("1", "系统架构及原型设计", "完成技术选型、模块划分、数据流设计\n系统接口定义、开发实施计划、产品原型验证",
     "架构", C_PRIMARY),
    ("2", "平台厂家协议沟通", "三地协议收集、技术方案分析\n厂家合作关系建立、协议测试工具开发",
     "协议", C_ACCENT),
]
for i, (num, title, desc, tag, color) in enumerate(core_items):
    x = Inches(0.6) + Inches(6.2) * i
    y = Inches(1.9)
    rrect(s, x, y, Inches(5.9), Inches(2.1), C_HL_BG, border=C_ACCENT)
    badge(s, x + Inches(0.2), y + Inches(0.2), Inches(0.5), Inches(0.35), num, color)
    txt(s, x + Inches(0.85), y + Inches(0.15), Inches(3.5), Inches(0.4),
        title, sz=18, color=C_DARK, bold=True)
    badge(s, x + Inches(4.7), y + Inches(0.18), Inches(0.9), Inches(0.3),
          "核心", C_ACCENT, sz=11)
    txt(s, x + Inches(0.25), y + Inches(0.7), Inches(5.4), Inches(1.2),
        desc, sz=14, color=C_GRAY)

# 其他任务
txt(s, Inches(0.6), Inches(4.2), Inches(6), Inches(0.42),
    "产品原型验证", sz=18, color=C_PRIMARY, bold=True)

other_items = [
    ("3", "播放引擎原型", "加载/播放/插值/设备输出全链路验证", C_PRIMARY),
    ("4", "动作编辑器原型", "时间轴编辑器窗口骨架与UI框架", C_PRIMARY),
    ("5", "影片播控平台方案", "影片播控平台需求分析与界面原型设计", C_TEAL),
    ("6", "动感文件格式规范", ".motion 文件标准 + JSON Schema", C_TEAL),
]
for i, (num, title, desc, color) in enumerate(other_items):
    col, row = i % 2, i // 2
    x = Inches(0.6) + Inches(6.2) * col
    y = Inches(4.7) + Inches(1.05) * row
    rrect(s, x, y, Inches(5.9), Inches(0.9), C_LIGHT_BG)
    badge(s, x + Inches(0.2), y + Inches(0.25), Inches(0.5), Inches(0.35),
          num, color)
    txt(s, x + Inches(0.85), y + Inches(0.1), Inches(2.5), Inches(0.35),
        title, sz=15, color=C_DARK, bold=True)
    txt(s, x + Inches(0.85), y + Inches(0.45), Inches(4.8), Inches(0.35),
        desc, sz=12, color=C_GRAY)

txt(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.4),
    "共 6 项计划目标，其中系统架构及原型设计与平台协议沟通为本阶段核心工作",
    sz=14, color=C_PRIMARY)


# ═══════════════════════════════════════════
# Slide 4: 已完成工作 (Change #3: 术语更新)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "二、已完成工作")

completed = [
    ["序号", "类别", "计划任务", "完成情况", "状态"],
    ["1", "核心", "系统架构及原型设计", "完整架构文档 v8，含技术选型、模块划分、数据流图、接口定义 + 实施计划", "已完成"],
    ["2", "核心", "平台厂家协议沟通", "邯郸/随州/温州三地协议收集 + 厂家合作关系建立 + 测试工具开发", "已完成"],
    ["3", "原型", "播放引擎原型", "加载/播放/插值/设备输出/停止回中，全流程验证通过", "已完成"],
    ["4", "原型", "动作编辑器原型", "编辑器独立窗口，工具栏/轨道面板/属性区布局就绪", "已完成"],
    ["5", "原型", "影片播控平台方案", "影片播控平台需求分析 + 界面原型设计（6 个核心界面）完成", "已完成"],
    ["6", "原型", "动感文件格式规范", ".motion JSON v1.0 定稿 + JSON Schema + 示例文件", "已完成"],
]
cws = [Inches(0.6), Inches(0.8), Inches(2.2), Inches(7.2), Inches(1.0)]
# 使用 tbl_ex 并对所有数据行(1-6)设置绿色已完成标识
completed_row_colors = {i: {'bg': C_OK_BG, 'text': C_OK_TEXT} for i in range(1, len(completed))}
ts_completed = tbl_ex(s, Inches(0.6), Inches(1.2), Inches(11.8), completed, cws,
                      row_colors=completed_row_colors)

# 完成率
rrect(s, Inches(0.6), Inches(4.6), Inches(5.5), Inches(0.6), C_OK_BG)
txt(s, Inches(1.0), Inches(4.65), Inches(5), Inches(0.5),
    "计划任务完成率：6 / 6（100%）", sz=18, color=C_TEAL, bold=True)

# 指标
txt(s, Inches(0.6), Inches(5.5), Inches(12), Inches(0.4),
    "关键成果指标", sz=16, color=C_DARK, bold=True)
metrics = [
    ("架构文档", "v8"), ("动作文件协议", "v1.0"),
    ("插值验证", "10000帧"), ("影院厂商协议沟通", "4家"),
    ("影片播控平台原型", "6页"),
]
for i, (label, value) in enumerate(metrics):
    x = Inches(0.6) + Inches(2.5) * i
    rrect(s, x, Inches(5.9), Inches(2.2), Inches(0.6), C_LIGHT_BG)
    txt(s, x, Inches(5.9), Inches(2.2), Inches(0.3),
        label, sz=11, color=C_GRAY, align=PP_ALIGN.CENTER)
    txt(s, x, Inches(6.15), Inches(2.2), Inches(0.3),
        value, sz=14, color=C_PRIMARY, bold=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════
# Slides 5-10: 影片播控平台界面原型 (Change #4: 前置, Change #5: 每页一张图)
# ═══════════════════════════════════════════
owner_slides_data = [
    (IMG_OW_STANDBY, "① 待机界面", "影厅待机状态，触摸屏幕开始选片"),
    (IMG_OW_SELECT,  "② 影片选择界面", "轮播选片 · 影片详情 · 推荐指数评级"),
    (IMG_OW_PLAY,    "③ 影片播放界面", "全屏播放 · 播放进度控制 · 动感平台同步"),
    (IMG_OW_EMERG,   "④ 紧急处理界面", "一键紧急停止 · 状态回显 · 恢复/终止操作"),
    (IMG_OW_END,     "⑤ 影片结束界面", "平台归位确认 · 重播/返回选片/待机"),
    (IMG_OW_MONITOR, "⑥ 现场监控界面", "多路摄像头 · 影厅/观众/设备实时画面"),
]

for idx, (img, title, desc) in enumerate(owner_slides_data):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C_DARK)
    txt(s, Inches(0.6), Inches(0.15), Inches(11), Inches(0.6),
        f"成果展示 · 影片播控平台 — {title}", sz=28, color=C_WHITE, bold=True)
    txt(s, Inches(0.6), Inches(0.62), Inches(11), Inches(0.35),
        desc, sz=14, color=C_LGRAY)
    # 单张大图居中
    add_img(s, img, Inches(1.65), Inches(1.05), w=Inches(10.0))
    # 最后一页加总结
    if idx == len(owner_slides_data) - 1:
        rrect(s, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.7),
              RGBColor(0x0E, 0x22, 0x3A))
        mtxt(s, Inches(1.0), Inches(6.65), Inches(11), Inches(0.6), [
            ("影片播控平台 — 6 个核心界面原型设计完成", True, C_ACCENT),
            ("覆盖完整放映流程：待机 → 选片 → 播放 → 监控 → 紧急处理 → 播放完成",
             False, C_LGRAY),
        ], sz=13)


# ═══════════════════════════════════════════
# Slide 11: 成果展示 — 动作编辑平台原型
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_DARK)
txt(s, Inches(0.6), Inches(0.2), Inches(10), Inches(0.7),
    "成果展示 · 动作编辑平台原型", sz=28, color=C_WHITE, bold=True)
add_img(s, IMG_EDITOR, Inches(0.4), Inches(1.1), w=Inches(12.5))
txt(s, Inches(0.6), Inches(6.8), Inches(12), Inches(0.5),
    "时间轴编辑器：视频预览区 + 关键帧属性面板 + 多轨道时间轴编辑区  |  姿态控制 (Yaw/Pitch/Roll) + 特效控制 (喷水/落雪/热风等)",
    sz=13, color=C_LGRAY)


# ═══════════════════════════════════════════
# Slide 12: 成果展示 — 飞荇客测试工具
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_DARK)
txt(s, Inches(0.6), Inches(0.2), Inches(10), Inches(0.7),
    "成果展示 · 飞荇客平台协议测试工具", sz=28, color=C_WHITE, bold=True)
add_img(s, IMG_TESTTOOL, Inches(2.8), Inches(1.0), w=Inches(7.7))
features = [
    "平台姿态控制 (Yaw/Pitch/Roll)",
    "特效控制 (喷水/雪/雾/热风/闪电/震动)",
    "播放控制 (5 路节目选择 + 停止)",
    "门角度监控 (4 门实时显示)",
    "压杆 / 门控 / 车辆控制",
    "通信日志实时记录",
]
for i, feat in enumerate(features):
    col, row = i % 3, i // 3
    x = Inches(0.6) + Inches(4.2) * col
    y = Inches(6.1) + Inches(0.4) * row
    txt(s, x, y, Inches(4.0), Inches(0.35), f"  {feat}", sz=12, color=C_LGRAY)


# ═══════════════════════════════════════════
# Slide 13: 厂家对接历程 (Change #7: Excel整合, 红绿标识, 对接人, 新总结)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "三、执行过程 · 厂家对接历程")

# 厂家对接数据: (日期, 事件, 成果, 对接人, color_type)
# color_type: "red"=不顺利, "green"=顺利, None=中性
timeline_events = [
    ("3/3",  "项目成立",
     "启动集成式动感平台播控系统开发", "杨天威", None),
    ("3/12", "收到邯郸影院动态平台控制协议",
     "确认厂家为北京和利时；厂家不愿提供更多信息", "邵向阳", "red"),
    ("3/13", "收到随州动感平台控制协议",
     "厂家仅提供平台 SDK，不提供姿态控制协议", "郝晓阳", "red"),
    ("3/16", "收到温州动感平台控制协议",
     "确认为北京和利时 6 轴直接控制平台", "穆大强", None),
    ("3/18", "取得北京和利时官方协议说明",
     "明确 UDP 电缸脉冲 + 体感座椅 SDK 双控制模式", "穆大强", "green"),
    ("3/19", "取得温州现场轨道影院可执行文件",
     "分析出基本技术架构", "马宝全", None),
    ("3/19", "取得邯郸流程控制简易通讯协议",
     "厂家建议购买测试设备，不要用现场设备测试", "马宝全", "red"),
    ("3/19", "取得邯郸影院可执行文件",
     "—", "邵向阳", None),
    ("3/20", "考虑就濉溪项目继续保持与厂家交涉",
     "—", "邵向阳/马宝全", None),
    ("3/20", "温州平台厂家明确不提供相关协议支持",
     "该协议对接方向终止", "—", "red"),
    ("3/23", "经友好沟通，邯郸厂家同意提供协议并配合协助",
     "建立合作关系，协议对接条件具备", "马宝全", "green"),
    ("3/24", "根据飞荇客协议进行测试工具开发",
     "测试工具开发完成，等待测试验证", "马宝全", "green"),
]

nr = len(timeline_events) + 1  # +1 for header
nc = 4
col_ws_tl = [Inches(0.7), Inches(4.2), Inches(4.5), Inches(1.8)]
ts = s.shapes.add_table(nr, nc, Inches(0.5), Inches(1.15),
                        Inches(12.2), Inches(0.38) * nr)
table = ts.table
for i, cw in enumerate(col_ws_tl):
    table.columns[i].width = cw

# 表头
for ci, h in enumerate(["日期", "事件", "成果", "对接人"]):
    cell = table.cell(0, ci); cell.text = ""
    p = cell.text_frame.paragraphs[0]; p.text = h; p.font.name = "微软雅黑"
    p.font.size = Pt(11); p.font.color.rgb = C_WHITE; p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    cell.fill.solid(); cell.fill.fore_color.rgb = C_TBL_HDR
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(4); cell.margin_right = Pt(4)
    cell.margin_top = Pt(3); cell.margin_bottom = Pt(3)

# 数据行
for ri, (date, event, result, person, ctype) in enumerate(timeline_events):
    row_idx = ri + 1
    if ctype == "red":
        bg_c, text_c = C_RED_BG, C_RED_TEXT
    elif ctype == "green":
        bg_c, text_c = C_OK_BG, C_GREEN_TEXT
    else:
        bg_c = C_TBL_R1 if row_idx % 2 == 1 else C_TBL_R2
        text_c = C_DARK
    for ci, ct in enumerate([date, event, result, person]):
        cell = table.cell(row_idx, ci); cell.text = ""
        p = cell.text_frame.paragraphs[0]; p.text = str(ct); p.font.name = "微软雅黑"
        p.font.size = Pt(9); p.font.color.rgb = text_c
        p.font.bold = (ctype is not None)
        cell.fill.solid(); cell.fill.fore_color.rgb = bg_c
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(4); cell.margin_right = Pt(4)
        cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)

# 链式节点总结
rrect(s, Inches(0.5), Inches(6.0), Inches(12.2), Inches(1.1), C_OK_BG)
mtxt(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(1.0), [
    ("历程总结：", True, C_DARK),
    ("三地协议收集 → 确认厂家为北京和利时 → 多次沟通建立合作关系 → 取得官方协议 → 测试工具开发", False, C_DARK),
    ("最终结果：平台测试工具开发完成，等待邯郸、濉溪项目进行协议验证", True, C_GREEN_TEXT),
], sz=12)


# ═══════════════════════════════════════════
# Slide 14: 开发工作推进 (Change #8: 重命名阶段, 多人并行)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "三、执行过程 · 开发工作推进")

phases = [
    ("第一周 3/9 - 3/14", "架构与原型开发", [
        "完成系统架构设计文档",
        "确定技术选型方案",
        "制定开发实施计划与工时估算",
        "影片播控平台需求分析",
    ], C_PRIMARY),
    ("第二周 3/15 - 3/21", "数据模型与文件格式", [
        "定义 .motion JSON 文件格式 v1.0",
        "编写 JSON Schema 验证规范",
        "实现核心数据模型层",
        "实现 4 种插值算法并验证精度",
    ], C_TEAL),
    ("第三周 3/22 - 3/28", "影院厂家平台协议对接", [
        "三地协议收集与分析",
        "邯郸厂家合作关系建立",
        "飞荇客协议测试工具开发",
        "影片播控平台界面原型设计",
    ], C_ACCENT),
]

for i, (week, title, items, color) in enumerate(phases):
    x = Inches(0.6) + Inches(4.2) * i
    y = Inches(1.3)
    rrect(s, x, y, Inches(3.9), Inches(4.5), C_LIGHT_BG)
    rect(s, x, y, Inches(3.9), Inches(0.65), color)
    txt(s, x + Inches(0.2), y + Inches(0.05), Inches(3.5), Inches(0.3),
        week, sz=12, color=C_WHITE, bold=True)
    txt(s, x + Inches(0.2), y + Inches(0.3), Inches(3.5), Inches(0.3),
        title, sz=15, color=C_WHITE, bold=True)
    for j, item in enumerate(items):
        txt(s, x + Inches(0.25), y + Inches(0.9) + Inches(0.6) * j,
            Inches(3.4), Inches(0.55), f"  {item}", sz=13, color=C_DARK)

for i in range(2):
    x = Inches(0.6) + Inches(3.9) * (i + 1) + Inches(0.3) * i
    txt(s, x - Inches(0.1), Inches(3.2), Inches(0.5), Inches(0.5),
        ">>", sz=28, color=C_ACCENT, bold=True, align=PP_ALIGN.CENTER)

txt(s, Inches(0.6), Inches(6.2), Inches(12), Inches(0.4),
    "多人并行推进：架构设计、原型开发、厂家协议对接同步进行",
    sz=14, color=C_PRIMARY)


# ═══════════════════════════════════════════
# Slide 15: 难点与突破 (Change #6: 移除难点三, 保留2个)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "四、难点与突破")

difficulties = [
    ("难点一：厂家协议获取困难",
     "问题",
     "影院厂家不愿开放相关协议，认为与其核心技术冲突\n"
     "不放心别人使用他们的平台，担心造成平台损坏等问题",
     "解决方案",
     "积极对接邯郸、温州、濉溪、随州等项目的轨道影院厂家，"
     "以互利共赢为目标，反复沟通，最终在邯郸、濉溪项目上取得突破，"
     "成功获取到对接协议",
     C_ORANGE),
    ("难点二：无设备条件下的协议验证",
     "问题",
     "取得的协议文档无法在公司进行验证",
     "解决方案",
     "根据协议开发测试工具，最大化的保障操作可控性，避免意外操作的发生\n"
     "积极分析邯郸项目现场的厂家软件，深度了解轨道影院的播控原理\n"
     "濉溪项目正在施工阶段，目前正在每天跟踪轨道影院的施工进度，"
     "在具备调试条件后，第一时间现场介入，对相关协议进行测试验证",
     C_ORANGE),
]

for i, (title, lbl1, prob, lbl2, sol, color) in enumerate(difficulties):
    x = Inches(0.6) + Inches(6.2) * i
    y = Inches(1.3)
    rrect(s, x, y, Inches(5.9), Inches(5.5), C_LIGHT_BG)
    rect(s, x, y, Inches(5.9), Inches(0.55), color)
    txt(s, x + Inches(0.2), y + Inches(0.08), Inches(5.5), Inches(0.4),
        title, sz=16, color=C_WHITE, bold=True)
    txt(s, x + Inches(0.2), y + Inches(0.7), Inches(5.5), Inches(0.3),
        lbl1, sz=14, color=C_ORANGE, bold=True)
    txt(s, x + Inches(0.2), y + Inches(1.1), Inches(5.5), Inches(1.2),
        prob, sz=13, color=C_GRAY)
    rect(s, x + Inches(0.2), y + Inches(2.5), Inches(5.5), Inches(0.02), C_LGRAY)
    txt(s, x + Inches(0.2), y + Inches(2.7), Inches(5.5), Inches(0.3),
        lbl2, sz=14, color=C_TEAL, bold=True)
    txt(s, x + Inches(0.2), y + Inches(3.1), Inches(5.5), Inches(2.2),
        sol, sz=13, color=C_DARK)


# ═══════════════════════════════════════════
# Slide 16: 厂家沟通实证
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "四、难点突破 · 厂家沟通实证")

# 左侧：3张微信聊天截图并排
txt(s, Inches(0.5), Inches(1.1), Inches(8), Inches(0.35),
    "厂家沟通记录", sz=14, color=C_DARK, bold=True)

add_img(s, IMG_CHAT1, Inches(0.5), Inches(1.5), w=Inches(2.5))
txt(s, Inches(0.5), Inches(5.7), Inches(2.5), Inches(0.6),
    "与飞荇客厂家沟通\n取得协议 + 说明书", sz=10, color=C_GRAY, align=PP_ALIGN.CENTER)

add_img(s, IMG_CHAT2, Inches(3.2), Inches(1.5), w=Inches(2.5))
txt(s, Inches(3.2), Inches(5.7), Inches(2.5), Inches(0.6),
    "反复沟通，成功获取\n关键端口信息", sz=10, color=C_GRAY, align=PP_ALIGN.CENTER)

add_img(s, IMG_CHAT3, Inches(5.9), Inches(1.5), w=Inches(2.5))
txt(s, Inches(5.9), Inches(5.7), Inches(2.5), Inches(0.6),
    "取得平台协议 + MDBOX 说明书\n9 分钟通话确认技术细节", sz=10, color=C_GRAY, align=PP_ALIGN.CENTER)

# 右侧：厂家对接时间线
txt(s, Inches(8.8), Inches(1.1), Inches(4), Inches(0.35),
    "厂家对接全过程", sz=14, color=C_DARK, bold=True)
add_img(s, IMG_TIMELINE, Inches(8.8), Inches(1.5), w=Inches(4.2))

# 底部总结
rrect(s, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.9), C_OK_BG)
mtxt(s, Inches(0.9), Inches(6.42), Inches(11.5), Inches(0.85), [
    ("结论：", True, C_DARK),
    ("以互利共赢为目标反复沟通 → 成功获取北京和利时官方协议 + MDBOX 说明书 + "
     "关键端口信息 → 测试工具开发完成，待设备到位后实机验证",
     False, C_DARK),
], sz=12)


# ═══════════════════════════════════════════
# Slide 17: 下一步计划 (Change #9: 影片播控平台首位, Change #2: 完整进度节点)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "五、下一步计划")

txt(s, Inches(0.6), Inches(1.15), Inches(12), Inches(0.4),
    f"{NODE_LABELS[2]}（2026 年 4 月 30 日前）：核心播控系统开发与基础联调完成",
    sz=18, color=C_DARK, bold=True)

next_items = [
    ("1", "影片播控平台开发",
     "影片播控平台核心开发\n视频播放管理\n动感平台联动控制", C_ACCENT),
    ("2", "动作编辑器功能开发",
     "时间轴编辑器完整实现\n关键帧编辑/曲线编辑/视频同步\n导出 .motion 文件", C_PRIMARY),
    ("3", "基础联调与协议验证",
     "编辑器与播放引擎对接\n飞荇客平台协议实机验证\n端到端流程打通", C_TEAL),
]

for i, (num, title, desc, color) in enumerate(next_items):
    x = Inches(0.6) + Inches(4.2) * i
    y = Inches(1.7)
    rrect(s, x, y, Inches(3.8), Inches(2.6), C_LIGHT_BG)
    badge(s, x + Inches(0.2), y + Inches(0.25), Inches(0.5), Inches(0.35),
          num, color, sz=16)
    txt(s, x + Inches(0.85), y + Inches(0.2), Inches(2.7), Inches(0.4),
        title, sz=18, color=C_DARK, bold=True)
    txt(s, x + Inches(0.3), y + Inches(0.9), Inches(3.2), Inches(1.5),
        desc, sz=14, color=C_GRAY)

# 完整进度节点表 (从 Slide 2 复制, Change #2)
plan_full = [
    ["节点", "截止日期", "进度安排（立项报告）", "状态"],
    [NODE_LABELS[1], "3/31", "产品原型设计及系统架构方案完成", "已按计划完成"],
    [NODE_LABELS[2], "4/30", "核心播控系统开发与基础联调完成", ""],
    [NODE_LABELS[3], "5/15", "完成设备接口整合及动作编排优化", ""],
    [NODE_LABELS[4], "5/31", "系统整体测试优化完成，达到上线标准", ""],
    [NODE_LABELS[5], "6/30", "产品验收完成", ""],
]
tbl_ex(s, Inches(0.6), Inches(4.6), Inches(12), plan_full,
       [Inches(2.0), Inches(1.2), Inches(6.5), Inches(2.3)],
       row_colors={1: {'bg': C_OK_BG, 'text': C_OK_TEXT}})


# ═══════════════════════════════════════════
# Slide 18: 人员分工与倒排计划 (Change #10: 节点一绿色, 厂家沟通工作)
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_WHITE)
title_bar(s, "六、人员分工与倒排计划")

# 倒排总览表
schedule_data = [
    ["节点", "截止", "马宝全", "居向前", "邵向阳", "郝晓阳", "穆大强"],
    [f"{NODE_LABELS[1]} ✅", "3/31",
     "架构+引擎+编辑器\n测试工具开发\n邯郸厂家沟通",
     "影片播控平台\n界面原型设计",
     "邯郸/濉溪\n厂家协议收集",
     "随州\n厂家协议收集",
     "温州/和利时\n厂家协议沟通"],
    [NODE_LABELS[2], "4/30",
     "编辑器完善\nTCP播控\n播放器核心",
     "影片播控平台\n架构+影片控制",
     "协议DLL架构\n基础通讯层", "—", "—"],
    [NODE_LABELS[3], "5/15",
     "C++ MotionPlayer\n编排优化",
     "平台控制\n播放器通信",
     "DLL核心封装\n指令处理",
     "Unity Package\n运行时组件",
     "UE插件\nFMotionClient"],
    [NODE_LABELS[4], "5/31",
     "效果预设\n性能优化",
     "影片播控平台联调",
     "协议联调\n异常处理",
     "编辑器扩展\nP/Invoke",
     "蓝图集成\n编辑器扩展"],
    [NODE_LABELS[5], "6/30",
     "授权系统\n打包+验收",
     "优化+部署",
     "文档+稳定性",
     "Demo+UPM发布",
     "Demo+uplugin"],
]
nr_s, nc_s = len(schedule_data), len(schedule_data[0])
cws_s = [Inches(2.0), Inches(0.7), Inches(2.0), Inches(1.8),
         Inches(1.8), Inches(1.6), Inches(1.8)]
ts = s.shapes.add_table(nr_s, nc_s, Inches(0.35), Inches(1.15),
                        Inches(12.6), Inches(0.55) * nr_s)
table = ts.table
for i, cw in enumerate(cws_s):
    table.columns[i].width = cw
for ri, row in enumerate(schedule_data):
    for ci, ct in enumerate(row):
        cell = table.cell(ri, ci)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        # 处理多行文本
        lines = str(ct).split('\n')
        p.text = lines[0]
        for line in lines[1:]:
            np = cell.text_frame.add_paragraph()
            np.text = line
            np.font.name = "微软雅黑"
            np.font.size = Pt(9)
            np.space_before = Pt(0)
            np.space_after = Pt(0)
            if ri == 0:
                np.font.color.rgb = C_WHITE
            elif ri == 1:
                np.font.color.rgb = C_GREEN_TEXT
            else:
                np.font.color.rgb = C_DARK
        p.font.name = "微软雅黑"
        if ri == 0:
            p.font.size = Pt(10)
            p.font.color.rgb = C_WHITE
            p.font.bold = True
            p.alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_TBL_HDR
        elif ri == 1:
            # 节点一已完成 — 绿色背景 + 绿色字体
            p.font.size = Pt(9)
            p.font.color.rgb = C_GREEN_TEXT
            p.font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_OK_BG
        else:
            p.font.size = Pt(9)
            p.font.color.rgb = C_DARK
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_TBL_R1 if ri % 2 == 0 else C_TBL_R2
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(4)
        cell.margin_right = Pt(4)
        cell.margin_top = Pt(2)
        cell.margin_bottom = Pt(2)

# 关键里程碑检查点
txt(s, Inches(0.6), Inches(5.1), Inches(12), Inches(0.35),
    "里程碑检查点", sz=14, color=C_DARK, bold=True)

checkpoints = [
    ("4/15", "编辑器核心可用 + 影片播控平台架构评审 + DLL框架评审"),
    ("4/30", "编辑器+TCP端到端 + 影片控制原型 + 通讯层心跳"),
    ("5/15", "C++ MotionPlayer驱动设备 + Unity/UE连接 + DLL指令链路"),
    ("5/31", "全系统联调 + Unity/UE编辑器扩展 + 协议联调"),
    ("6/15", "Demo场景完成 + 文档初稿 + 安装包初版"),
    ("6/30", "产品验收：全部交付物就绪"),
]
for i, (date, desc) in enumerate(checkpoints):
    col, row = i % 3, i // 3
    x = Inches(0.6) + Inches(4.2) * col
    y = Inches(5.5) + Inches(0.85) * row
    rrect(s, x, y, Inches(3.9), Inches(0.72), C_LIGHT_BG)
    badge(s, x + Inches(0.1), y + Inches(0.17), Inches(0.65), Inches(0.32),
          date, C_ACCENT, sz=10)
    txt(s, x + Inches(0.85), y + Inches(0.07), Inches(2.9), Inches(0.6),
        desc, sz=10, color=C_DARK)


# ═══════════════════════════════════════════
# Slide 19: 结束页
# ═══════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, C_DARK)
rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.06), C_ACCENT)
txt(s, Inches(0), Inches(2.0), SLIDE_W, Inches(1),
    "谢谢", sz=56, color=C_WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0), Inches(3.3), SLIDE_W, Inches(0.6),
    "集成式动感平台播控系统  |  第一阶段汇报",
    sz=22, color=C_LGRAY, align=PP_ALIGN.CENTER)
txt(s, Inches(0), Inches(4.5), SLIDE_W, Inches(0.5),
    "2026 年 3 月 31 日", sz=18, color=C_LGRAY, align=PP_ALIGN.CENTER)
rect(s, Inches(0), Inches(7.3), SLIDE_W, Inches(0.04), C_ACCENT)


# ═══════════════════════════════════════════
prs.save(OUTPUT_PATH)
print(f"PPT generated: {OUTPUT_PATH}")
